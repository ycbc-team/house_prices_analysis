#!/usr/bin/env python3
"""
东风东教育集团番禺小学 教师招聘数据分析脚本
从 raw/ 目录下的政府公示Excel中提取数据，筛选并统计
"""

import openpyxl
import csv
import os
from collections import Counter

RAW_DIR = os.path.join(os.path.dirname(__file__), '..', 'raw')
PARSED_DIR = os.path.join(os.path.dirname(__file__), '..', 'parsed')

# 所有Excel文件配置
FILES = [
    # 优才计划
    ('youcai_2024s1.xlsx', '2024年优才计划第一批'),
    ('youcai_2024b1.xlsx', '2024年教育系统优才计划第一批'),
    ('youcai_2025h1.xlsx', '2025年上半年优才计划'),
    ('youcai_2025.xlsx', '2025年优才计划'),
    ('youcai_2025h2.xlsx', '2025年下半年优才计划'),
    # 统考
    ('tongkao_2025b1.xlsx', '2025年统考第一批'),
    ('tongkao_2025b2.xlsx', '2025年统考第二批'),
    ('tongkao_2026b1.xlsx', '2026年统考第一批'),
]

TARGET_KEYWORDS = ['东风东', '番禺小学']


def find_header(ws):
    """找到表头行和列索引"""
    for r in range(1, min(6, ws.max_row + 1)):
        row_vals = [str(c.value) if c.value else '' for c in ws[r]]
        if '招聘单位' in row_vals:
            headers = row_vals
            return r, headers
    return None, None


def parse_excel(filepath, label):
    """解析单个Excel文件"""
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    header_row, headers = find_header(ws)
    if not header_row:
        return []

    # 找列索引
    col_map = {}
    for i, h in enumerate(headers):
        h_clean = h.strip().replace(' ', '')
        if h_clean in ['招聘单位', '姓名', '学历', '最高学历', '学位', '所学专业', '专业', '毕业院校', '招聘岗位', '岗位']:
            col_map[h_clean] = i

    results = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        unit = str(row[col_map.get('招聘单位', 0)]) if col_map.get('招聘单位') is not None else ''

        # 检查是否匹配目标学校
        is_target = any(kw in unit for kw in TARGET_KEYWORDS)

        record = {
            'batch': label,
            'unit': unit,
            'post': row[col_map.get('招聘岗位', col_map.get('岗位', 3))] if col_map.get('招聘岗位') or col_map.get('岗位') else '',
            'name': row[col_map.get('姓名', 4)] if col_map.get('姓名') else '',
            'edu': row[col_map.get('最高学历', col_map.get('学历', 6))] if col_map.get('最高学历') or col_map.get('学历') else '',
            'degree': row[col_map.get('学位', 7)] if col_map.get('学位') else '',
            'major': row[col_map.get('所学专业', col_map.get('专业', 8))] if col_map.get('所学专业') or col_map.get('专业') else '',
            'school': row[col_map.get('毕业院校', 9)] if col_map.get('毕业院校') else '',
            'is_target': is_target,
        }
        results.append(record)

    return results


def main():
    os.makedirs(PARSED_DIR, exist_ok=True)

    all_records = []
    target_records = []
    unit_stats = {}

    for fname, label in FILES:
        fpath = os.path.join(RAW_DIR, fname)
        if not os.path.exists(fpath):
            print(f"跳过（不存在）: {fname}")
            continue

        records = parse_excel(fpath, label)
        all_records.extend(records)

        # 统计每个批次的单位分布
        units = Counter([r['unit'] for r in records if r['unit'] and r['unit'] != 'None'])
        unit_stats[label] = dict(units)

        # 筛选目标学校
        targets = [r for r in records if r['is_target']]
        target_records.extend(targets)

        print(f"{label}: {len(records)}人, 目标学校匹配: {len(targets)}人")

    # 导出所有记录为CSV
    all_csv = os.path.join(PARSED_DIR, 'all_teachers.csv')
    with open(all_csv, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=['batch', 'unit', 'post', 'name', 'edu', 'degree', 'major', 'school', 'is_target'])
        writer.writeheader()
        writer.writerows(all_records)

    # 导出目标学校记录
    target_csv = os.path.join(PARSED_DIR, 'dongfengdong_teachers.csv')
    with open(target_csv, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=['batch', 'unit', 'post', 'name', 'edu', 'degree', 'major', 'school'])
        writer.writeheader()
        writer.writerows([{k: v for k, v in r.items() if k != 'is_target'} for r in target_records])

    # 导出单位统计
    stats_csv = os.path.join(PARSED_DIR, 'unit_distribution.csv')
    with open(stats_csv, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(['批次', '招聘单位', '人数'])
        for batch, units in unit_stats.items():
            for unit, count in sorted(units.items(), key=lambda x: -x[1]):
                writer.writerow([batch, unit, count])

    print(f"\n解析完成：")
    print(f"  总记录数: {len(all_records)}")
    print(f"  东风东番禺小学匹配: {len(target_records)}")
    print(f"  输出文件:")
    print(f"    - {all_csv}")
    print(f"    - {target_csv}")
    print(f"    - {stats_csv}")


if __name__ == '__main__':
    main()
