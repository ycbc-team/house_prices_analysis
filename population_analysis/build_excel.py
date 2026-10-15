# -*- coding: utf-8 -*-
"""从 raw CSV 生成汇总 Excel 工作簿。"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(BASE, 'raw')
OUT = os.path.join(BASE, '广州各区人口变化_数据集.xlsx')

HDR_FILL = PatternFill('solid', fgColor='1F4E79')
HDR_FONT = Font(color='FFFFFF', bold=True, size=11)
SEC_FILL = PatternFill('solid', fgColor='D9E2F3')
SEC_FONT = Font(bold=True, size=11, color='1F4E79')
BORDER = Border(*[Side(style='thin', color='BBBBBB')]*4)

def load_rows(fn):
    with open(os.path.join(RAW, fn), encoding='utf-8-sig') as f:
        return list(csv.reader(f))

def long_to_wide(rows, year_col=0, dist_col=1, val_col=2, drop_note=True):
    """长表 -> 宽表 {年份: {区: 值}}"""
    w = {}
    for r in rows:
        if r[0] == '年份':
            continue
        try:
            y = int(r[year_col])
        except Exception:
            continue
        w.setdefault(y, {})[r[dist_col]] = float(r[val_col])
    return w

def write_wide_sheet(wb, name, title, wide, order=None, note='单位：万人'):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = Font(bold=True, size=14, color='1F4E79')
    ws['A2'] = note; ws['A2'].font = Font(size=10, color='666666')
    years = sorted(wide)
    dists = order or []
    if not dists:
        seen = set()
        for y in years:
            for d in wide[y]:
                if d not in seen:
                    seen.add(d); dists.append(d)
    # 区名顺序: 现11区优先，旧区跟随
    prio = ['荔湾区','越秀区','海珠区','天河区','白云区','黄埔区','番禺区','花都区','南沙区','从化区','增城区',
            '东山区','芳村区','萝岗区','增城市','从化市']
    dists = [d for d in prio if any(d in wide[y] for y in years)] + \
            [d for d in dists if d not in prio]
    ws.cell(3, 1, '年份')
    for c, d in enumerate(dists, start=2):
        cell = ws.cell(3, c, d)
        cell.fill = HDR_FILL; cell.font = HDR_FONT; cell.alignment = Alignment(horizontal='center')
    for c, y in enumerate(years, start=1):
        ws.cell(c + 3, 1, y).font = Font(bold=True)
        for cc, d in enumerate(dists, start=2):
            v = wide[y].get(d)
            cell = ws.cell(c + 3, cc, v if v is not None else '')
            cell.border = BORDER
            if v is not None:
                cell.number_format = '0.00'
    ws.column_dimensions['A'].width = 8
    for c in range(2, len(dists) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.freeze_panes = 'B4'
    ws.sheet_view.showGridLines = False

def write_long_sheet(wb, name, title, rows, widths=None):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = Font(bold=True, size=14, color='1F4E79')
    hdr = rows[0]
    for c, h in enumerate(hdr, start=1):
        cell = ws.cell(3, c, h)
        cell.fill = HDR_FILL; cell.font = HDR_FONT; cell.alignment = Alignment(horizontal='center')
    for r_i, r in enumerate(rows[1:], start=4):
        for c, v in enumerate(r, start=1):
            cell = ws.cell(r_i, c, v)
            cell.border = BORDER
    if widths:
        for c, w_ in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w_
    ws.freeze_panes = 'A4'
    ws.sheet_view.showGridLines = False

wb = Workbook()
wb.remove(wb.active)

# 说明页
ws = wb.create_sheet('说明')
ws['A1'] = '广州各区人口变化数据集（1998—2025）'
ws['A1'].font = Font(bold=True, size=16, color='1F4E79')
notes = [
    '',
    '一、数据来源（全部为广州市统计局官方公开资料）',
    '  1. 历年《广州统计年鉴》（1999—2025 版）第二章“人口”篇及第一章区县主要指标；',
    '  2. 广州市历年《人口规模及分布情况》公报（2024 年末、2025 年末两版）；',
    '  3. 第五/六/七次全国人口普查公报及《广州市区域人口分布和人口流动情况分析》（七普专项，含 2010/2020 分区跨市流入与净流入）。',
    '',
    '二、口径说明',
    '  1. 户籍人口：1998—2024 来自统计年鉴“年末户籍总人口”，2024/2025 来自公报表 2。',
    '  2. 常住人口：2005—2014 为统计年鉴估计口径（2023 版 2-4 表）；2015—2023 为按七普修订口径（2024 版 2-2 表）；',
    '     2024/2025 为公报口径；2000/2010/2020 另有普查标准时点（11 月 1 日）数据，与年末数不可直接混用。',
    '  3. 自然变动（出生/死亡）与迁移（迁入/迁出）均为户籍口径，单位为人、‰。',
    '  4. 区划沿革断档（拼接时必须分段）：',
    '     · 2005 年前：东山区、芳村区存在（2005 年分别并入越秀区、荔湾区）；',
    '     · 2005—2014：增设南沙区、萝岗区（2014-02 萝岗并入黄埔区）；',
    '     · 2012：南沙区接收番禺区大岗/东涌/榄核三镇（番禺、南沙数据有断档）；',
    '     · 2015 起：11 区口径（荔湾/越秀/海珠/天河/白云/黄埔/番禺/花都/南沙/从化/增城）。',
    '  5. 已知缺口：2004—2006 年分区自然/迁移明细无公开数据（对应版年鉴人口篇文件加密），仅存 2004 年分区户籍；',
    '     2025 年分区自然/迁移明细未发布（公报仅全市合计数）。相关年份在表中留空。',
    '  6. “跨市流入”为普查口径（居住本区且户口登记地在外市），2010/2020 两个时点；“净流入”= 常住人口-户籍人口。',
    '',
    '三、校验结果',
    '  户籍人口：1998—2025 各年 12/11 区分区合计与年鉴/公报全市数完全一致；',
    '  常住人口：年度口径合计与年鉴、公报全市数一致；普查时点合计与普查公报一致；',
    '  自然/迁移：1998—2003、2007—2024 分区合计与全市数一致（2024 与 2025 版年鉴锚点核对一致）。',
    '  详细逐条校验输出见 build_dataset.py 运行日志。',
    '',
    '四、文件结构',
    '  分区户籍人口_1998_2025.csv / 分区常住人口_2000_2025.csv / 分区自然变动_1998_2025.csv /',
    '  分区户籍迁移_1998_2025.csv / 普查跨市流入与净流入_2010_2020.csv / 全市长序列_1978_2025.csv',
]
for i, t in enumerate(notes, start=2):
    cell = ws.cell(i, 1, t)
    if t.startswith(('一、', '二、', '三、', '四、')):
        cell.font = SEC_FONT
ws.column_dimensions['A'].width = 110
ws.sheet_view.showGridLines = False

# 户籍宽表
hk = load_rows('分区户籍人口_1998_2025.csv')
write_wide_sheet(wb, '户籍人口_分区', '分区户籍人口（万人，1998—2025）', long_to_wide(hk),
                 note='单位：万人。2005 前含东山区/芳村区，2005—2014 含南沙区/萝岗区，2015 起 11 区。详见“说明”页区划断档。')
# 常住宽表
cz = load_rows('分区常住人口_2000_2025.csv')
write_wide_sheet(wb, '常住人口_分区', '分区常住人口（万人，2000—2025）', long_to_wide(cz),
                 note='单位：万人。2005—2014 年鉴估计口径；2015—2023 七普修订口径；2000/2010/2020 普查时点值单独标出（详见说明）。')
# 自然变动长表
nat = load_rows('分区自然变动_1998_2025.csv')
write_long_sheet(wb, '自然变动_分区', '分区户籍自然变动（人、‰，1998—2024）', nat,
                 widths=[8, 12, 14, 10, 10, 10, 10, 12])
# 迁移长表
mig = load_rows('分区户籍迁移_1998_2025.csv')
write_long_sheet(wb, '户籍迁移_分区', '分区户籍迁入迁出（人、‰，1998—2024）', mig,
                 widths=[8, 12, 10, 10, 10, 10, 10, 10])
# 跨市流入
ld = load_rows('普查跨市流入与净流入_2010_2020.csv')
write_long_sheet(wb, '跨市流入_普查', '分区跨市流入与净流入（万人，2010/2020 普查时点）', ld,
                 widths=[12, 16, 16, 14, 14])
# 全市长序列
city = load_rows('全市长序列_1978_2025.csv')
write_long_sheet(wb, '全市长序列', '全市户籍人口/常住人口/自然变动/迁移（1978—2025）', city,
                 widths=[8, 20, 14, 24])

wb.save(OUT)
print('EXCEL OK:', OUT)
